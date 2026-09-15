"""Excel 导出：全部/勾选/搜索结果快照；电话按文本保存，保留 +86 与前导零。"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADERS = ["姓名", "公司", "职位", "电话1", "电话2", "邮箱", "地址", "业务", "结识展会",
           "结识时间", "备注", "印刷标签", "推断标签", "状态"]
KEYS = ["name", "company", "title", "phone1", "phone2", "email", "address", "business",
        "event", "met_at", "notes", "tags_printed", "tags_inferred", "status"]
TEXT_COLS = {"电话1", "电话2"}  # 文本格式列


def export_xlsx(contacts: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "联系人"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for c in contacts:
        row = []
        for k in KEYS:
            v = c.get(k, "")
            if isinstance(v, list):
                v = "、".join(v)
            row.append(v)
        ws.append(row)
    # 电话列强制文本格式，保留 +86 / 前导零
    for idx, h in enumerate(HEADERS, start=1):
        if h in TEXT_COLS:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=idx).number_format = "@"
    widths = [12, 22, 16, 18, 18, 24, 26, 22, 16, 14, 26, 18, 18, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
